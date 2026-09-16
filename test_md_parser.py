import re
import docx
from docx.shared import Pt, RGBColor

def format_cell_markdown(p, text, base_font="Calibri", base_size=10):
    p.text = "" # clear existing text
    
    # 1. Pre-process LaTeX and math symbols to clean academic typography
    replacements = [
        (r"$\mathcal{O}(n \cdot \log n)$", "O(n · log n)"),
        (r"$\mathcal{O}(n^2)$", "O(n²)"),
        (r"$\mathcal{O}(n^3)$", "O(n³)"),
        (r"$\mathcal{O} (n \cdot k \cdot i) $", "O(n · k · i)"),
        (r"$R = \frac{P_{unique}}{N_{pkts}}$", "R = P_unique / N_pkts"),
        (r"$\le 4$", "≤ 4"),
        (r"$\sim 60$", "≈ 60"),
        (r"$\to$", "→"),
        (r"$s \approx 0.44$", "s ≈ 0.44"),
        (r"$s = 0.88$", "s = 0.88"),
        (r"$s = 0.92$", "s = 0.92"),
        (r"$R > 0.9$", "R > 0.9"),
        (r"$\mu_L = \frac{1}{N} \sum Length$", "μ_L = (1/N) ∑ Length"),
        (r"$N_{pkts} = \sum 1$", "N_pkts = ∑ 1"),
        (r"$N_{SYN} = \sum SYN\_Flag$", "N_SYN = ∑ SYN_Flag"),
        (r"$P_{unique} = \bigcup Dst\_Port$", "P_unique = |⋃ Dst_Port|"),
        (r"$IP_{unique} = \bigcup Dst\_Address$", "IP_unique = |⋃ Dst_Address|"),
    ]
    for old, new in replacements:
        text = text.replace(old, new)
        
    # Check if the entire cell is wrapped in **...**
    is_entirely_bold = False
    if text.startswith("**") and text.endswith("**") and text.count("**") == 2:
        is_entirely_bold = True
        text = text[2:-2]

    # Pattern for remaining tokens: **bold**, *italic*, `code`
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            content = part[2:-2]
            run = p.add_run(content)
            run.bold = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
        elif part.startswith('*') and part.endswith('*'):
            content = part[1:-1]
            run = p.add_run(content)
            run.italic = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
        elif part.startswith('`') and part.endswith('`'):
            content = part[1:-1]
            run = p.add_run(content)
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 0.5)
            if is_entirely_bold:
                run.bold = True
        else:
            run = p.add_run(part)
            if is_entirely_bold:
                run.bold = True
            run.font.name = base_font
            run.font.size = Pt(base_size)

doc = docx.Document()
test_cases = [
    "**Mécanisme fondamental**",
    "**Nombre d'arbres (`n_estimators`)**",
    "$\\mathcal{O}(n \\cdot \\log n)$ • *Ultra-rapide*",
    "Chaque paquet SYN a une taille légitime ($\\sim 60$ B) $\\to$ **Non détecté**",
    "Traitement des flux en direct sur interface physique ou virtuelle (`raw sockets`) et analyse différée de fichiers de capture (`.pcap`, `.pcapng`).",
    "Définition d'un ratio d'entropie $R = \\frac{P_{unique}}{N_{pkts}}$ combiné au décompte des drapeaux `SYN` pour discriminer la navigation Web d'un balayage hostile."
]

for tc in test_cases:
    p = doc.add_paragraph()
    format_cell_markdown(p, tc)

doc.save("test_md_out.docx")
print("All advanced test cases rendered successfully!")


