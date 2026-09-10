import sys
import docx
from docx.shared import Cm

sys.stdout.reconfigure(encoding='utf-8')
doc_path = r"C:\Users\DevLab\Desktop\Soutenance\Avec Commune\Rapport_MohamedRedaQIYAOUI.docx"
doc = docx.Document(doc_path)

# Ensure margins are set on all sections
for s in doc.sections:
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(2.5)

doc.save(doc_path)

print(f"Total sections: {len(doc.sections)}")
for i, s in enumerate(doc.sections):
    h_text = " ".join(p.text for p in s.header.paragraphs if p.text.strip())
    f_text = " ".join(p.text for p in s.footer.paragraphs if p.text.strip())
    print(f"Section {i}: margins=({s.top_margin.cm:.1f}, {s.bottom_margin.cm:.1f}, {s.left_margin.cm:.1f}, {s.right_margin.cm:.1f}) | Header: '{h_text}' | Footer: '{f_text}'")

total_words = sum(len(p.text.split()) for p in doc.paragraphs)
print(f"Total paragraphs: {len(doc.paragraphs)}, Total words: {total_words}")

img_count = sum(1 for rel in doc.part.rels.values() if "image" in rel.target_ref)
print(f"Total embedded images in docx: {img_count}")
print(f"Total tables: {len(doc.tables)}")

print("\n--- Check Cover Page ---")
for i in range(26):
    txt = doc.paragraphs[i].text.strip()
    if txt:
        print(f"  P{i:02d}: {txt}")

print("\n--- Check Annexes at end ---")
for p in doc.paragraphs[-10:]:
    txt = p.text.strip()
    if txt:
        print(f"  [{p.style.name}]: {txt[:80]}")
