import docx

doc = docx.Document(r"C:\Users\DevLab\Desktop\Rapport_Soutenance_QIYAOUI.docx")

print("=== COVER PAGE RUNS (0 to 32) ===")
for i in range(min(32, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip():
        runs = [(r.text.strip(), r.font.name, r.font.size.pt if r.font.size else None, str(r.font.color.rgb) if r.font.color else None, r.font.bold) for r in p.runs if r.text.strip()]
        print(f"P{i:02d} [align={p.alignment}]: '{p.text.strip()}' | runs={runs}")
