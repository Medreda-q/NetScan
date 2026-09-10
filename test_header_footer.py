import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Section 0: Preliminary
p1 = doc.add_paragraph("Couverture / Titre")
p2 = doc.add_paragraph("Sommaire...")
pPr = p2._p.get_or_add_pPr()
sectPr = parse_xml(f'<w:sectPr {nsdecls("w")}><w:type w:val="nextPage"/></w:sectPr>')
pPr.append(sectPr)

# Section 1: Body
p3 = doc.add_paragraph("Introduction Générale...")
p4 = doc.add_paragraph("Corps du rapport...")

sec_prelim = doc.sections[0]
sec_body = doc.sections[1]

# Set body header
sec_body.header.is_linked_to_previous = False
hp = sec_body.header.paragraphs[0]
hp.text = "EST Tétouan — Cybersécurité\t\tNetScan : Détection d'Anomalies Réseau"
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
for r in hp.runs:
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = RGBColor(100, 100, 100)

# Add bottom border to header paragraph
pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="4" w:color="0A2540"/></w:pBdr>')
hp._p.get_or_add_pPr().append(pBdr)

# Set body footer with dynamic page number
sec_body.footer.is_linked_to_previous = False
fp = sec_body.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

# Restart page numbering at 1
sec_body_sectPr = sec_body._sectPr
pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>')
sec_body_sectPr.append(pgNumType)

# Add dynamic page field
r_page = fp.add_run()
r_page.font.name = "Times New Roman"
r_page.font.size = Pt(10)
r_page.font.color.rgb = RGBColor(100, 100, 100)

fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
r_page._r.append(fldChar1)
r_page._r.append(instrText)
r_page._r.append(fldChar2)
r_page._r.append(fldChar3)

doc.save("test_output_sections.docx")
print("Saved test_output_sections.docx successfully!")
