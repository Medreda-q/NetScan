import docx, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

doc = docx.Document(r"C:\Users\DevLab\Desktop\Soutenance\Rapport_MohamedRedaQIYAOUI_PERFECT.docx")

print("=== SCANNING PARAGRAPHS FOR LITERAL MARKDOWN / LATEX ===")
p_leaks = 0
for i, p in enumerate(doc.paragraphs):
    txt = p.text
    has_md_bold = "**" in txt
    has_code = "`" in txt
    has_dollar = "$" in txt
    has_latex = any(cmd in txt for cmd in [r"\to", r"\implies", r"\frac", r"\ln", r"\mathcal", r"\bigcup", r"\subset", r"\sim", r"\ge", r"\le", r"\lceil", r"\lfloor"])
    has_raw_ast = "*" in txt and not txt.strip().startswith("•") and not txt.strip().startswith("* ")
    if has_md_bold or has_code or has_dollar or has_latex or has_raw_ast:
        print(f"P[{i}] (Style: {p.style.name}): {txt}")
        p_leaks += 1

print(f"Total paragraph leaks: {p_leaks}")

print("\n=== SCANNING TABLES FOR LITERAL MARKDOWN / LATEX ===")
t_leaks = 0
for t_idx, tbl in enumerate(doc.tables):
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            txt = cell.text
            has_md_bold = "**" in txt
            has_code = "`" in txt
            has_dollar = "$" in txt
            has_latex = any(cmd in txt for cmd in [r"\to", r"\implies", r"\frac", r"\ln", r"\mathcal", r"\bigcup", r"\subset", r"\sim", r"\ge", r"\le", r"\lceil", r"\lfloor"])
            has_raw_ast = "*" in txt and not txt.strip().startswith("•")
            if has_md_bold or has_code or has_dollar or has_latex or has_raw_ast:
                print(f"Table[{t_idx}] Row[{r_idx}] Col[{c_idx}]: {txt}")
                t_leaks += 1

print(f"Total table leaks: {t_leaks}")



