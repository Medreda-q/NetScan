import sys
import os
import shutil
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

backup_path = r"C:\Users\DevLab\Desktop\Soutenance\Avec Commune\Rapport_MohamedRedaQIYAOUI_BACKUP.docx"
target_path = r"C:\Users\DevLab\Desktop\Soutenance\Avec Commune\Rapport_MohamedRedaQIYAOUI.docx"

print(f"Loading document from backup: {backup_path}")
doc = docx.Document(backup_path)

# 1. Margins: 2.5 cm on all sections
for i, section in enumerate(doc.sections):
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
print("[+] Step 1: Margins set to 2.5 cm on all sections.")

# 2. Page de garde
for p in doc.paragraphs[:24]:
    for r in p.runs:
        r.font.name = "Times New Roman"

idx_author = -1
for i in range(24):
    if "Mohamed Reda Qiyaoui" in doc.paragraphs[i].text:
        idx_author = i
        break

if idx_author != -1:
    p_author = doc.paragraphs[idx_author]
    
    p_org = p_author.insert_paragraph_before()
    p_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_org = p_org.add_run("Organisme d'accueil : Commune de Tétouan\nService du SIG, des Systèmes Informatiques et de la Numérisation")
    r_org.font.name = "Times New Roman"
    r_org.font.size = Pt(11)
    r_org.bold = True
    r_org.font.color.rgb = RGBColor(10, 37, 64)
    p_org.paragraph_format.space_before = Pt(8)
    p_org.paragraph_format.space_after = Pt(8)

    p_enc = p_author.insert_paragraph_before()
    p_enc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r1 = p_enc.add_run("Encadrant Pédagogique : ")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.bold = True
    r2 = p_enc.add_run("Corps Enseignant — EST Tétouan\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)

    r3 = p_enc.add_run("Encadrant Professionnel : ")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(11)
    r3.bold = True
    r4 = p_enc.add_run("Responsable Informatique & SIG — Commune de Tétouan")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(11)
    
    p_enc.paragraph_format.space_before = Pt(4)
    p_enc.paragraph_format.space_after = Pt(12)
    print("[+] Step 2: Page de garde enriched with Organisme d'accueil and Encadrants.")

# 3. Section break right before Introduction Générale
idx_intro = -1
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "Introduction Générale" and p.style.name.startswith("Heading"):
        idx_intro = i
        break

if idx_intro != -1:
    p_before_intro = doc.paragraphs[idx_intro - 1]
    pPr = p_before_intro._p.get_or_add_pPr()
    sectPr = parse_xml(f'<w:sectPr {nsdecls("w")}><w:type w:val="nextPage"/></w:sectPr>')
    pPr.append(sectPr)
    print(f"[+] Inserted section break right before Introduction Générale.")

doc.save(target_path)
doc = docx.Document(target_path)

# Header & Footer
doc.sections[0].different_first_page_header_footer = True

if len(doc.sections) > 1:
    doc.sections[1].header.is_linked_to_previous = False
    for p in doc.sections[1].header.paragraphs:
        p.text = ""
    doc.sections[1].footer.is_linked_to_previous = False
    for p in doc.sections[1].footer.paragraphs:
        p.text = ""

if len(doc.sections) > 2:
    sec_body = doc.sections[2]
    sec_body.header.is_linked_to_previous = False
    hp = sec_body.header.paragraphs[0]
    hp.text = "EST Tétouan — Filière Cybersécurité\t\tNetScan : Détection d'Anomalies Réseau"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in hp.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(8.5)
        r.font.italic = True
        r.font.color.rgb = RGBColor(110, 110, 110)
    
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="4" w:color="0A2540"/></w:pBdr>')
    hp._p.get_or_add_pPr().append(pBdr)

    sec_body.footer.is_linked_to_previous = False
    fp = sec_body.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    sec_body_sectPr = sec_body._sectPr
    for existing in sec_body_sectPr.xpath('./w:pgNumType'):
        sec_body_sectPr.remove(existing)
    pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="1"/>')
    sec_body_sectPr.append(pgNumType)

    r_lbl = fp.add_run("Page ")
    r_lbl.font.name = "Times New Roman"
    r_lbl.font.size = Pt(10)
    r_lbl.font.color.rgb = RGBColor(100, 100, 100)

    r_num = fp.add_run()
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(10)
    r_num.bold = True
    r_num.font.color.rgb = RGBColor(10, 37, 64)

    fld1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fld2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fld3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r_num._r.append(fld1)
    r_num._r.append(instr)
    r_num._r.append(fld2)
    r_num._r.append(fld3)
    print("[+] Step 3: Headers & Footers configured with page numbering starting at 1 at Introduction Générale.")

# 4. Typography & 1.5 line spacing
body_idx = -1
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "Introduction Générale" and p.style.name.startswith("Heading"):
        body_idx = i
        break

for i, p in enumerate(doc.paragraphs):
    is_body = (i >= body_idx) if body_idx != -1 else (i >= 24)
    sn = p.style.name

    if sn.startswith("Heading 1"):
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(16)
            r.bold = True
            r.font.color.rgb = RGBColor(10, 37, 64)
    elif sn.startswith("Heading 2"):
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(14)
            r.bold = True
            r.font.color.rgb = RGBColor(24, 76, 120)
    elif sn.startswith("Heading 3"):
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = RGBColor(30, 41, 59)
    else:
        if is_body:
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.name = "Times New Roman"
            if not r.font.size:
                r.font.size = Pt(12)
            elif is_body and sn == "Normal" and r.font.size.pt < 11.5:
                r.font.size = Pt(12)

print("[+] Step 4: Times New Roman 12 pt & 1.5 line spacing applied to body text.")

# 5. Insert Figures with robust paragraph detection
fig_syn_scan_path = r"c:\Users\DevLab\Documents\GitHub\EST\Soutenance Project\schema_nmap_syn_scan_fr.png"
fig_syn_flood_path = r"c:\Users\DevLab\Documents\GitHub\EST\Soutenance Project\schema_syn_flood_fr.png"
fig_ratio_path = r"c:\Users\DevLab\Documents\GitHub\EST\Soutenance Project\presentation_assets\chart_port_diversity_ratio.png"
fig_benchmark_path = r"c:\Users\DevLab\Documents\GitHub\EST\Soutenance Project\presentation_assets\chart_benchmark_metrics.png"

# Figure 2.1: Before 2.2.2
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("2.2.2") and os.path.exists(fig_syn_scan_path):
        p_img = p.insert_paragraph_before()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(fig_syn_scan_path, width=Inches(5.7))

        p_cap = p.insert_paragraph_before()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        rc = p_cap.add_run("Figure 2.1 : Fonctionnement du Three-Way Handshake comparé au scan furtif TCP SYN semi-ouvert (Nmap -sS).")
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(10)
        rc.font.italic = True
        rc.font.color.rgb = RGBColor(60, 60, 60)
        print("[+] Step 5a: Figure 2.1 (TCP SYN Stealth Scan) inserted.")
        break

# Figure 2.2: Before 2.3
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("2.3 ") and os.path.exists(fig_syn_flood_path):
        p_img = p.insert_paragraph_before()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(fig_syn_flood_path, width=Inches(5.7))

        p_cap = p.insert_paragraph_before()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        rc = p_cap.add_run("Figure 2.2 : Schéma temporel de l'attaque TCP SYN Flood et saturation de la Backlog Queue.")
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(10)
        rc.font.italic = True
        rc.font.color.rgb = RGBColor(60, 60, 60)
        print("[+] Step 5b: Figure 2.2 (TCP SYN Flood) inserted.")
        break

# Figure 4.3: Port Diversity Ratio
for i, p in enumerate(doc.paragraphs):
    if ("Matrice de dispersion comportementale" in p.text or "unique_dst_ports / packet_count" in p.text) and os.path.exists(fig_ratio_path):
        p_img = p.insert_paragraph_before()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(fig_ratio_path, width=Inches(5.5))

        p_cap = p.insert_paragraph_before()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        rc = p_cap.add_run("Figure 4.3 : Matrice de dispersion comportementale (Diversité de ports uniques vs Volume de paquets).")
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(10)
        rc.font.italic = True
        rc.font.color.rgb = RGBColor(60, 60, 60)
        print("[+] Step 5c: Figure 4.3 (Port Diversity Scatter) inserted.")
        break

# Figure 5.1: Benchmark Metrics
for i, p in enumerate(doc.paragraphs):
    if "Analyse des résultats et validation expérimentale" in p.text and os.path.exists(fig_benchmark_path):
        p_target = doc.paragraphs[i + 2]
        p_img = p_target.insert_paragraph_before()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(fig_benchmark_path, width=Inches(5.5))

        p_cap = p_target.insert_paragraph_before()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        rc = p_cap.add_run("Figure 5.1 : Métriques comparatives de performance et validation opérationnelle (Précision, Rappel, F1-Score).")
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(10)
        rc.font.italic = True
        rc.font.color.rgb = RGBColor(60, 60, 60)
        print("[+] Step 5d: Figure 5.1 (Benchmark Metrics) inserted.")
        break

# 6. Add Annexes
doc.add_page_break()
p_annexe_h = doc.add_paragraph("ANNEXES", style="Heading 1")

p_annexe_a = doc.add_paragraph("Annexe A : Fiche Technique d'Architecture & Modules Logiciels NetScan", style="Heading 2")
p_annexe_intro = doc.add_paragraph(
    "Le tableau ci-dessous synthétise la cartographie modulaire de NetScan telle que présentée aux membres du jury. "
    "Il détaille le rôle opérationnel de chaque fichier, ses flux d'entrée/sortie et les commandes d'exécution correspondantes."
)
p_annexe_intro.paragraph_format.line_spacing = 1.5

table_data = [
    ("Module / Fichier", "Rôle & Mission Principale", "Entrées / Sources", "Sorties / Livrables", "Commande d'Exécution"),
    ("netscan.py", "Orchestrateur principal & CLI", "Arguments CLI, flux réseau brut", "Logs console ANSI, initialisation", "python netscan.py --help"),
    ("detection_engine.py", "Moteur IA Dual-Layer (Isolation Forest)", "Flux direct (Npcap/libpcap) ou PCAP", "Scores d'anomalie [0.0 - 1.0], alertes", "python netscan.py --pcap <capture.pcap>"),
    ("dashboard.py", "Tableau de bord SOC interactif", "Métriques streaming, logs CSV", "Interface visuelle Streamlit temps réel", "python netscan.py --dashboard"),
    ("alerting.py", "Alerting instantané & défense active", "Événements anormaux (score >= 0.60)", "Webhooks Discord (< 300 ms), règles iptables", "Automatique en cours d'analyse"),
    ("report_generator.py", "Générateur exécutif de rapports forensiques", "Métriques d'incident, timeline", "Rapports HTML5 et PDF (NIST SP 800-61)", "python netscan.py --pcap <file> --report"),
    ("test_smoke.py", "Banc de tests et assurance qualité", "Environnement local Python", "Validation des 5 modules en < 3 secondes", "python test_smoke.py"),
]

table = doc.add_table(rows=len(table_data), cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
col_widths = [Cm(2.6), Cm(4.1), Cm(3.2), Cm(3.2), Cm(3.4)]

for r_idx, row in enumerate(table_data):
    for c_idx, val in enumerate(row):
        cell = table.cell(r_idx, c_idx)
        cell.width = col_widths[c_idx]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        
        run = p.add_run(val)
        run.font.name = "Times New Roman"
        if r_idx == 0:
            run.font.size = Pt(9.5)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0A2540"/>')
            cell._tc.get_or_add_tcPr().append(shading)
        else:
            run.font.size = Pt(8.5)
            if c_idx == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(10, 37, 64)
            if r_idx % 2 == 1:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
                cell._tc.get_or_add_tcPr().append(shading)

tblPr = table._tbl.tblPr
tblBorders = parse_xml(f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        <w:insideV w:val="none"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
    </w:tblBorders>
''')
tblPr.append(tblBorders)

p_tbl_cap = doc.add_paragraph("Tableau A.1 : Matrice descriptive des composants logiciels du système NetScan.", style="Normal")
p_tbl_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_tbl_cap.paragraph_format.space_after = Pt(14)
for r in p_tbl_cap.runs:
    r.font.name = "Times New Roman"
    r.font.size = Pt(9.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor(90, 90, 90)

p_annexe_b = doc.add_paragraph("Annexe B : Guide d'Exécution et Traces Opérationnelles", style="Heading 2")
p_annexe_b_txt = doc.add_paragraph(
    "NetScan prend en charge deux modes d'exploitation distincts, adaptés respectivement à l'investigation forensique "
    "post-mortem et à la détection proactive en environnement de production :\n\n"
    "1. Mode Forensique Hors-Ligne (Rejeu de capture PCAP) :\n"
    "   python netscan.py --pcap sample_nmap_attack.pcap --report\n"
    "   Ce mode lit séquentiellement les trames, extrait les vecteurs 5D, applique la prédiction Isolation Forest "
    "et génère un rapport forensique d'incident complet conforme aux directives NIST SP 800-61.\n\n"
    "2. Mode Surveillance Temps Réel (Calibrate-Then-Monitor) :\n"
    "   python netscan.py --live --interface eth0 --calibrate 120 --window 5\n"
    "   Ce mode observe le trafic normal pendant 120 secondes pour calibrer la ligne de base, puis bascule en surveillance "
    "continue par micro-fenêtres de 5 secondes avec notification Discord et génération automatique de règles de blocage iptables."
)
p_annexe_b_txt.paragraph_format.line_spacing = 1.5
for r in p_annexe_b_txt.runs:
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)

print("[+] Step 6: Annexes added.")

# Save final document
doc.save(target_path)
print(f"[SUCCESS] Complete updated report saved to:\n  {target_path}")
