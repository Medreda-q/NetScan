"""
apply_qiyaoui_design.py — Builds the final docx with 100% fidelity to the visual identity of Rapport_Soutenance_QIYAOUI.docx.
Exact colors: #1F4E5C (Heading 1 / Primary), #2E86AB (Heading 2 / Table Header), #333333 (Heading 3), #555555 (Subtitles/Captions), #EAF3F5 (Table zebra fill).
Typography: Times New Roman throughout, Cambria Math, Consolas for code.
Includes official EST banner (image1.jpeg) and exact layout.
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="6", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_custom_report(md_path, docx_path, logo_path=None):
    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    doc = Document()

    # Define color palette matching Rapport_Soutenance_QIYAOUI.docx
    PRIMARY_TEAL   = RGBColor(0x1F, 0x4E, 0x5C)  # #1F4E5C
    OCEAN_BLUE     = RGBColor(0x2E, 0x86, 0xAB)  # #2E86AB
    DARK_GRAY      = RGBColor(0x33, 0x33, 0x33)  # #333333
    MEDIUM_GRAY    = RGBColor(0x55, 0x55, 0x55)  # #555555
    BODY_BLACK     = RGBColor(0x11, 0x11, 0x11)  # #111111

    # Section 0 setup: Cover Page margins matching reference (56.7pt top/bottom, 65.7pt left, 42.5pt right)
    sec0 = doc.sections[0]
    sec0.top_margin    = Pt(56.7)
    sec0.bottom_margin = Pt(56.7)
    sec0.left_margin   = Pt(65.7)
    sec0.right_margin  = Pt(42.5)
    sec0.page_width    = Pt(595.25)
    sec0.page_height   = Pt(841.85)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = BODY_BLACK
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Insert official header logo if available
    if logo_path and os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(0)
        p_logo.paragraph_format.space_after = Pt(14)
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(6.2))

    lines = md_content.split("\n")
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []
    first_page_break_handled = False

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows_data = []
        for l in tbl_lines:
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', l):
                continue
            cells = [c.strip() for c in l.strip().strip('|').split('|')]
            if cells and any(cells):
                rows_data.append(cells)
        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        for r in rows_data:
            while len(r) < num_cols:
                r.append("")

        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="2E86AB", sz="6")

        for row_idx, row_data in enumerate(rows_data):
            row = table.rows[row_idx]
            is_header = (row_idx == 0)
            for col_idx, cell_text in enumerate(row_data):
                cell = row.cells[col_idx]
                cell.text = cell_text
                set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05

                if is_header:
                    set_cell_background(cell, "2E86AB")  # Exact #2E86AB from reference
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(10)
                else:
                    bg = "EAF3F5" if row_idx % 2 == 0 else "FFFFFF"  # Exact #EAF3F5 zebra tint
                    set_cell_background(cell, bg)
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)
                        run.font.color.rgb = DARK_GRAY

        doc.add_paragraph()

    def flush_code(c_lines):
        if not c_lines:
            return
        code_text = "\n".join(c_lines)
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F3F4F6")
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        set_table_borders(tbl, color="D1D5DB", sz="6")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(code_text)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        doc.add_paragraph()

    i = 0
    while i < len(lines):
        line = lines[i]

        # Handle page breaks & transition to Section 1 (Body margins: 72pt top/bottom, 54pt left/right)
        if "\\newpage" in line:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

            if not first_page_break_handled:
                sec1 = doc.add_section()
                sec1.top_margin    = Pt(72.0)
                sec1.bottom_margin = Pt(72.0)
                sec1.left_margin   = Pt(54.0)
                sec1.right_margin  = Pt(54.0)
                sec1.page_width    = Pt(595.3)
                sec1.page_height   = Pt(841.9)
                first_page_break_handled = True
            else:
                doc.add_page_break()
            i += 1
            continue

        if line.strip() == "---":
            # Just separator
            i += 1
            continue

        # Code blocks
        if line.strip().startswith("```"):
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

            if in_code_block:
                flush_code(code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Tables
        if "|" in line and (line.strip().startswith("|") or line.strip().endswith("|")):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        # Blockquotes / Dédicaces
        if stripped.startswith(">"):
            quote_text = re.sub(r"^>\s*", "", stripped).replace("*", "").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(f"« {quote_text} »" if not quote_text.startswith("«") else quote_text)
            run.font.name = "Times New Roman"
            run.font.italic = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = MEDIUM_GRAY
            i += 1
            continue

        # Headings (Matching exact fonts & colors from reference)
        if stripped.startswith("# "):
            h_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            
            # Center if cover or main chapter
            if not first_page_break_handled or any(kw in h_text.upper() for kw in ["RAPPORT", "CONCEPTION", "REMERCIEMENTS", "RÉSUMÉ", "ABSTRACT", "TABLE DES MATIÈRES"]):
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.size = Pt(15)
            run.font.color.rgb = PRIMARY_TEAL
            i += 1
            continue

        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = OCEAN_BLUE
            i += 1
            continue

        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = DARK_GRAY
            i += 1
            continue

        if stripped.startswith("#### "):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = MEDIUM_GRAY
            i += 1
            continue

        # Bullet lists
        if stripped.startswith("- ") or stripped.startswith("* "):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.font.name = "Times New Roman"
                    run.font.bold = True
                    run.font.color.rgb = DARK_GRAY
                else:
                    run = p.add_run(part)
                    run.font.name = "Times New Roman"
                    run.font.color.rgb = BODY_BLACK
            i += 1
            continue

        # Numbered lists
        if re.match(r'^\d+\.\s', stripped):
            item_text = re.sub(r'^\d+\.\s*', '', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.font.name = "Times New Roman"
                    run.font.bold = True
                    run.font.color.rgb = DARK_GRAY
                else:
                    run = p.add_run(part)
                    run.font.name = "Times New Roman"
                    run.font.color.rgb = BODY_BLACK
            i += 1
            continue

        # Captions for figures/tables (e.g. *Tableau 1.1 : ...* or *Figure 2.1 : ...*)
        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[1:-1])
            run.font.name = "Times New Roman"
            run.font.italic = True
            run.font.size = Pt(10.0)
            run.font.color.rgb = MEDIUM_GRAY
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15

        # Cover page centering detection
        if not first_page_break_handled:
            if any(kw in stripped for kw in ["Mohamed Reda", "Année Universitaire", "Filière", "Présenté par", "Encadrement"]):
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Parse inline bold, code, math
        parts = re.split(r'(\*\*.*?\*\*|`.*?`|\$.*?\$)', stripped)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.font.name = "Times New Roman"
                run.font.bold = True
                if not first_page_break_handled:
                    run.font.size = Pt(12)
                    run.font.color.rgb = MEDIUM_GRAY
                else:
                    run.font.color.rgb = DARK_GRAY
            elif part.startswith("`") and part.endswith("`"):
                run = p.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)
            elif part.startswith("$") and part.endswith("$"):
                run = p.add_run(part[1:-1])
                run.font.name = "Cambria Math"
                run.font.italic = True
            else:
                run = p.add_run(part)
                run.font.name = "Times New Roman"
                run.font.color.rgb = BODY_BLACK

        i += 1

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    doc.save(docx_path)
    print(f"SUCCESS: Custom report perfectly styled and saved to: {docx_path}")

if __name__ == "__main__":
    logo = r"scratch/word/media/image1.jpeg"
    build_custom_report("RAPPORT_DE_STAGE_NETSCAN.md", "RAPPORT_DE_STAGE_NETSCAN.docx", logo_path=logo)
