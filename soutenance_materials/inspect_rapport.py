import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')
doc_path = r"C:\Users\DevLab\Desktop\Soutenance\Avec Commune\Rapport_MohamedRedaQIYAOUI.docx"
doc = docx.Document(doc_path)

print(f"Sections: {len(doc.sections)}")
for idx, sec in enumerate(doc.sections):
    print(f"Section {idx}: top={sec.top_margin.cm:.2f}cm, bottom={sec.bottom_margin.cm:.2f}cm, left={sec.left_margin.cm:.2f}cm, right={sec.right_margin.cm:.2f}cm")
    h_text = " ".join(p.text for p in sec.header.paragraphs if p.text.strip())
    f_text = " ".join(p.text for p in sec.footer.paragraphs if p.text.strip())
    print(f"  Header: '{h_text}'")
    print(f"  Footer: '{f_text}'")

print(f"\nTotal paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n--- Outline Sample ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    if p.style.name.startswith("Heading") or any(k in txt.upper() for k in ["CHAPITRE", "INTRODUCTION", "CONCLUSION", "SOMMAIRE", "REMERCIEMENTS"]):
        print(f"P{i:03d} [{p.style.name}] : {txt[:90]}")
