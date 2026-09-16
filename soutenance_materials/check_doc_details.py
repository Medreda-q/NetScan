import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')
doc_path = r"C:\Users\DevLab\Desktop\Soutenance\Avec Commune\Rapport_MohamedRedaQIYAOUI.docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
words = sum(len(p.text.split()) for p in doc.paragraphs)
chars = sum(len(p.text) for p in doc.paragraphs)
print(f"Total words: {words}")
print(f"Total characters: {chars}")

print("\n--- First 30 paragraphs (Page de garde & liminaires) ---")
for i in range(min(30, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"P{i:02d} [{p.style.name}]: {p.text}")

print("\n--- Sections & Margins ---")
for i, s in enumerate(doc.sections):
    print(f"Section {i}: top={s.top_margin.cm:.2f}cm, bottom={s.bottom_margin.cm:.2f}cm, left={s.left_margin.cm:.2f}cm, right={s.right_margin.cm:.2f}cm, orientation={s.orientation}")
    print(f"  header p count: {len(s.header.paragraphs)}, text: {[p.text for p in s.header.paragraphs if p.text.strip()]}")
    print(f"  footer p count: {len(s.footer.paragraphs)}, text: {[p.text for p in s.footer.paragraphs if p.text.strip()]}")

print("\n--- Paragraph Styles Used ---")
styles_used = set(p.style.name for p in doc.paragraphs)
print(styles_used)

print("\n--- Font details for sample paragraphs ---")
for p in doc.paragraphs[100:105]:
    fonts = [r.font.name for r in p.runs if r.font.name]
    sizes = [r.font.size.pt for r in p.runs if r.font.size]
    line_spacing = p.paragraph_format.line_spacing
    print(f"P: {p.text[:50]}... | Fonts: {set(fonts)} | Sizes: {set(sizes)} | Spacing: {line_spacing}")
