import re
import docx

def format_paragraph_markdown(p, base_font="Calibri", base_size=11):
    text = p.text
    if "*" not in text and "`" not in text and "$" not in text:
        return
    
    # Pre-clean known strings
    text = text.replace("$R$", "R").replace("$< 30$", "< 30")
    
    # Tokenize: **bold**, *italic*, `code`, text
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    parts = pattern.split(text)
    if len(parts) <= 1 and text == p.text:
        return
        
    p.text = "" # clear
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            run = p.add_run(part)

print("Paragraph formatter ready!")
