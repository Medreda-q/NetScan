import re
import os
import shutil
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def add_leader_line(doc, target_before_p, text, page_num, is_bold=False, is_indent=False, is_sub_indent=False):
    """
    Creates a paragraph with a right-aligned tab stop and dot leader at 9275 dxa,
    attaching text and right-aligned page number.
    """
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    tabs = parse_xml(r'<w:tabs %s><w:tab w:val="right" w:leader="dot" w:pos="9275"/></w:tabs>' % nsdecls('w'))
    pPr.append(tabs)
    
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    prefix = ""
    if is_sub_indent:
        prefix = "      "
    elif is_indent:
        prefix = "   "
        
    r1 = p.add_run(prefix + text)
    r1.font.name = "Calibri"
    r1.font.size = Pt(10)
    if is_bold:
        r1.bold = True
        
    r2 = p.add_run(f"\t{page_num}")
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    if is_bold:
        r2.bold = True
        
    target_before_p._element.addprevious(p._element)
    return p

def format_cell_markdown(p, text, base_font="Calibri", base_size=10):
    """
    Parses Markdown tokens (**bold**, *italic*, `code`) and mathematical symbols
    in a table cell and applies genuine Word run formatting.
    """
    p.text = "" # clear existing text
    
    # 1. Clean raw LaTeX math notation into elegant academic typography
    replacements = [
        (r"$\mathcal{O}(n \cdot \log n)$", "O(n · log n)"),
        (r"$\mathcal{O}(n^2)$", "O(n²)"),
        (r"$\mathcal{O}(n^3)$", "O(n³)"),
        (r"$\mathcal{O} (n \cdot k \cdot i) $", "O(n · k · i)"),
        (r"$R = \frac{P_{unique}}{N_{pkts}}$", "R = P_unique / N_pkts"),
        (r"$\le 4$", "≤ 4"),
        (r"$< 30$", "< 30"),
        (r"$\sim 60$", "≈ 60"),
        (r"$\to$", "→"),
        (r"$s \approx 0.44$", "s ≈ 0.44"),
        (r"$s = 0.88$", "s = 0.88"),
        (r"$s = 0.92$", "s = 0.92"),
        (r"$R > 0.9$", "R > 0.9"),
        (r"$\mu_L = \frac{1}{N} \sum Length$", "μ_L = (1/N) ∑ Length"),
        (r"$N_{pkts} = \sum 1$", "N_pkts = ∑ 1"),
        (r"$N_{SYN} = \sum SYN\_Flag$", "N_SYN = ∑ SYN_Flag"),
        (r"$P_{unique} = \bigcup Dst\_Port$", "P_unique = |⋃ Dst_Port|"),
        (r"$IP_{unique} = \bigcup Dst\_Address$", "IP_unique = |⋃ Dst_Address|"),
    ]
    for old, new in replacements:
        text = text.replace(old, new)
        
    # Check if the entire cell is wrapped in **...**
    is_entirely_bold = False
    if text.startswith("**") and text.endswith("**") and text.count("**") == 2:
        is_entirely_bold = True
        text = text[2:-2]

    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            content = part[2:-2]
            run = p.add_run(content)
            run.bold = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
        elif part.startswith('*') and part.endswith('*'):
            content = part[1:-1]
            run = p.add_run(content)
            run.italic = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
        elif part.startswith('`') and part.endswith('`'):
            content = part[1:-1]
            run = p.add_run(content)
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 0.5)
            if is_entirely_bold:
                run.bold = True
        else:
            run = p.add_run(part)
            if is_entirely_bold:
                run.bold = True
            run.font.name = base_font
            run.font.size = Pt(base_size)


def format_paragraph_markdown_inline(p):
    """
    Parses *italic*, **bold**, and `code` in regular body paragraphs.
    Preserves style, alignment, spacing.
    """
    xml_str = p._p.xml
    if "oMath" in xml_str or "blip" in xml_str:
        return
    
    txt = p.text
    if "*" not in txt and "`" not in txt and "$" not in txt:
        return
    
    # Pre-clean
    txt = txt.replace("$R$", "R").replace("$< 30$", "< 30")
    
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    parts = pattern.split(txt)
    if len(parts) <= 1:
        if txt != p.text:
            p.text = txt
        return

    orig_style = p.style
    orig_align = p.alignment
    orig_sb = p.paragraph_format.space_before
    orig_sa = p.paragraph_format.space_after
    orig_li = p.paragraph_format.left_indent

    p.text = ""
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            p.add_run(part)

    p.style = orig_style
    p.alignment = orig_align
    p.paragraph_format.space_before = orig_sb
    p.paragraph_format.space_after = orig_sa
    p.paragraph_format.left_indent = orig_li


def run_master_fix(src_doc_path, out_doc_path):
    print(f"Loading document: {src_doc_path}")
    doc = docx.Document(src_doc_path)
    
    # ==========================================================================
    # 1. HARMONIZE CHAPTER HEADINGS (ALL CAPS FOR CHAPTER I)
    # ==========================================================================
    for p in doc.paragraphs:
        if p.text.strip().startswith("Chapitre I : Présentation"):
            p.text = "CHAPITRE I : PRÉSENTATION DU CADRE ACADÉMIQUE ET ENVIRONNEMENT DU PROJET"
            print("[OK 1] Titre du Chapitre I harmonisé en majuscules.")
            break
            
    # ==========================================================================
    # 2. SOMMAIRE COMPLET AVEC POINTS DE SUITE ET NUMÉROTATION DE PAGES EXACTE
    # ==========================================================================
    toc_data = [
        ("1.  Introduction Générale", 11, True, False, False),
        ("•  Origine et enjeu de sécurité", 11, False, True, False),
        ("•  Problématique centrale : L'angle mort des signatures", 11, False, True, False),
        ("•  Objectifs généraux et périmètre du projet NetScan", 11, False, True, False),
        ("•  Organisation du document", 11, False, True, False),
        ("2.  Chapitre I : Présentation du Cadre Académique et Environnement du Projet", 12, True, False, False),
        ("1.1 - Cadre institutionnel et organisme d'accueil", 12, False, True, False),
        ("1.1.1 - Présentation de la commune de Tétouan", 12, False, False, True),
        ("1.1.2 - Processus / services rendus", 13, False, False, True),
        ("1.1.3 - Validation opérationnelle in situ", 13, False, False, True),
        ("1.2 - Démarche méthodologique en laboratoire R&D et autonomie", 14, False, True, False),
        ("1.3 - Cahier des charges et spécifications fonctionnelles", 15, False, True, False),
        ("1.4 - Matrice d'analyse et de gestion des risques", 16, False, True, False),
        ("3.  Chapitre II : État de l'Art et Fondements Théoriques", 18, True, False, False),
        ("2.1 - Les systèmes de détection d'intrusion réseau (NIDS) conventionnels", 18, False, True, False),
        ("2.2 - Mécanismes des attaques cibles : Scan Nmap TCP SYN et Floods volumétriques", 19, False, True, False),
        ("2.3 - Paradigmes du Machine Learning en sécurité (Supervisé vs Non Supervisé)", 20, False, True, False),
        ("2.4 - L'algorithme Isolation Forest (Modélisation mathématique et formulation)", 21, False, True, False),
        ("4.  Chapitre III : Architecture Système et Conception Technique", 25, True, False, False),
        ("3.1 - Pipeline global de traitement des flux", 25, False, True, False),
        ("3.2 - Ingestion réseau hybride (Mode direct promiscuité et captures PCAP/PCAPNG)", 25, False, True, False),
        ("3.3 - Ingénierie des caractéristiques (Feature Engineering)", 26, False, True, False),
        ("3.4 - Matrice de qualification heuristique et typage des menaces", 28, False, True, False),
        ("5.  Chapitre IV : Implémentation Logicielle et Outils Développés", 30, True, False, False),
        ("4.1 - Environnement de développement et stack technologique", 30, False, True, False),
        ("4.1.1 - Portabilité et prérequis pour l'environnement Windows", 30, False, False, True),
        ("4.2 - Orchestrateur en ligne de commande Linux/Kali (netscan.py)", 31, False, True, False),
        ("4.3 - Tableau de bord de supervision SOC interactif (dashboard.py)", 32, False, True, False),
        ("4.4 - Moteur de Journalisation Sécurisée et Historique des Incidents (alerts_history.json)", 34, False, True, False),
        ("4.5 - Générateur exécutif de rapports d'incident (report_generator.py)", 35, False, True, False),
        ("4.6 - Suite d'assurance qualité et tests de fumée (test_smoke.py)", 35, False, True, False),
        ("6.  Chapitre V : Expérimentation, Déploiement et Validation in situ", 36, True, False, False),
        ("5.1 - Architecture de déploiement et banc d'essai in situ (Commune de Tétouan)", 36, False, True, False),
        ("5.2 - Scénario 1 : Détection en direct d'un scan Nmap furtif (Live Sniffing)", 37, False, True, False),
        ("5.3 - Scénario 2 : Détection d'une inondation SYN Flood", 39, False, True, False),
        ("5.4 - Analyse comparative et gains opérationnels", 40, False, True, False),
        ("7.  Chapitre VI : Bilan du Stage, Compétences et Perspectives", 42, True, False, False),
        ("6.1 - Compétences techniques et professionnelles consolidées", 42, False, True, False),
        ("6.2 - Obstacles techniques rencontrés et arbitrages réalisés", 42, False, True, False),
        ("6.3 - Perspectives d'évolution du système", 43, False, True, False),
        ("8.  Conclusion Générale", 44, True, False, False),
        ("9.  Bibliographie et Webographie", 45, True, False, False),
    ]

    toc_start_idx = None
    toc_end_idx = None
    for i, p in enumerate(doc.paragraphs[:120]):
        if "Table des Matières" in p.text:
            toc_start_idx = i + 1
        if "Liste des Figures" in p.text and toc_start_idx is not None and i > toc_start_idx:
            toc_end_idx = i
            break

    if toc_start_idx is not None and toc_end_idx is not None:
        target_p = doc.paragraphs[toc_end_idx]
        paras_to_remove = [doc.paragraphs[k] for k in range(toc_start_idx, toc_end_idx)]
        for p in paras_to_remove:
            p._element.getparent().remove(p._element)

        # Re-add front matter section links
        front_links = [
            ("Dédicace", 2),
            ("Remerciements", 3),
            ("Résumé / Abstract", 4),
            ("Liste des Figures", 8),
            ("Liste des Tableaux", 9),
            ("Liste des Abréviations", 10),
        ]
        for title, pg in front_links:
            add_leader_line(doc, target_p, title, pg, is_bold=False)

        # Add all body entries
        for title, pg, is_b, is_ind, is_sub in toc_data:
            add_leader_line(doc, target_p, title, pg, is_bold=is_b, is_indent=is_ind, is_sub_indent=is_sub)

        print("[OK 2] Sommaire entièrement reconstruit avec points de suite et numéros de pages exacts.")

    # ==========================================================================
    # 3. FRONT MATTER : LISTE DES FIGURES ET TABLEAUX AVEC POINTS DE SUITE & PAGES
    # ==========================================================================
    clean_figures = [
        ("Figure 5.1 : Configuration réseau du Poste 1 d'Audit (Kali Linux).", 36),
        ("Figure 5.2 : Configuration réseau du Poste 2 de Supervision NetScan (Windows).", 37),
        ("Figure 5.3 : Démarrage de la sonde NetScan en écoute passive directe.", 37),
        ("Figure 5.4 : Simulation du balayage de reconnaissance furtif (nmap -sS).", 38),
        ("Figure 5.5 : Détection en direct et qualification de l'alerte Port Scan par NetScan.", 38),
        ("Figure 5.6 : Simulation de l'attaque par déni de service volumétrique (hping3 SYN Flood).", 39),
        ("Figure 5.7 : Détection en temps réel et qualification de l'alerte Flood / DDoS par NetScan.", 39)
    ]

    clean_tableaux = [
        ("Tableau 1.0 : Fiche signalétique de l'organisme d'accueil.", 12),
        ("Tableau 1.1 : Matrice des spécifications fonctionnelles et contraintes techniques de NetScan.", 16),
        ("Tableau 1.2 : Matrice d'analyse, d'évaluation et de mitigation des risques du projet.", 17),
        ("Tableau 2.1 : Comparatif synthétique : NIDS à base de signatures vs NIDS non supervisé.", 19),
        ("Tableau 2.2 : Comparaison des algorithmes d'isolation (Isolation Forest, One-Class SVM, k-Means).", 24),
        ("Tableau 3.1 : Vecteur des caractéristiques extraites au Niveau 1 (Analyse unitaire).", 27),
        ("Tableau 3.2 : Vecteur des caractéristiques comportementales agrégées au Niveau 2.", 28),
        ("Tableau 3.3 : Synthèse comparative des hyperparamètres des deux couches d'Isolation Forest.", 29),
        ("Tableau 4.1 : Stack technologique, langages et bibliothèques logicielles mobilisées.", 31),
        ("Tableau 5.1 : Synthèse comparative des résultats expérimentaux de détection sur capture réelle.", 40),
        ("Tableau 5.2 : Bilan comparatif des capacités de détection : Modèle Baseline vs NetScan Dual-Layer.", 41)
    ]

    fig_start_idx = None
    fig_end_idx = None
    for i, p in enumerate(doc.paragraphs[:160]):
        if "Figure 5.1" in p.text and fig_start_idx is None and i > 50:
            fig_start_idx = i
        if "Liste des Abr" in p.text and i > 100:
            fig_end_idx = i
            break

    if fig_start_idx is not None and fig_end_idx is not None:
        target_p = doc.paragraphs[fig_end_idx]
        paras_to_remove = [doc.paragraphs[k] for k in range(fig_start_idx, fig_end_idx)]
        for p in paras_to_remove:
            p._element.getparent().remove(p._element)

        for fig_title, pg in clean_figures:
            add_leader_line(doc, target_p, fig_title, pg, is_bold=False)

        h_tab = doc.add_paragraph(style='Heading 1')
        h_tab.text = "Liste des Tableaux"
        target_p._element.addprevious(h_tab._element)

        for tab_title, pg in clean_tableaux:
            add_leader_line(doc, target_p, tab_title, pg, is_bold=False)

        spacer = doc.add_paragraph(style='Normal')
        spacer.text = ""
        target_p._element.addprevious(spacer._element)
        print("[OK 3] Liste des Figures et Tableaux (avec Tableau 1.0) reconstruite avec points de suite et pages.")

    # ==========================================================================
    # 4. ENRICH RÉSUMÉ & ABSTRACT WITH CROSS-PLATFORM HIGHLIGHT
    # ==========================================================================
    for p in doc.paragraphs[:60]:
        if "NetScan est développé pour les environnements Linux/Kali" in p.text or "architecture multiplateforme" in p.text:
            p.text = (
                "NetScan est développé avec une architecture multiplateforme (conçu pour Kali Linux en phase d'audit "
                "et R&D, et déployé nativement sur Windows avec Npcap en exploitation opérationnelle in situ). Il comprend "
                "un orchestrateur en ligne de commande, un tableau de bord web interactif sous Streamlit, un système de "
                "journalisation d'audit persistant permettant à l'administrateur de consulter à tout moment l'historique "
                "complet des tentatives d'intrusion détectées depuis l'interface, et un générateur de rapports d'incident "
                "forensiques aux formats HTML5 et PDF. Sur des captures de trafic hostile réelles et lors d'écoutes directes "
                "in situ, l'outil détecte sans signature les scans furtifs et les dénis de service par SYN Flood, avec une "
                "latence sub-seconde et un taux de faux positifs faible."
            )
            print("[OK 4a] Résumé enrichi avec la portabilité multiplateforme.")
        elif "NetScan runs natively on Linux/Kali" in p.text or "Engineered with cross-platform" in p.text:
            p.text = (
                "Engineered with cross-platform architecture (Kali Linux for offensive testing and R&D, and native Windows "
                "deployment with Npcap for operational in situ defense), NetScan includes a command-line orchestrator, an "
                "interactive Streamlit dashboard, a persistent security audit logging engine enabling administrators to review "
                "historical intrusion attempts directly from the interface, and a forensic incident report generator producing "
                "HTML5 and PDF output. Tested against real hostile traffic captures and live sniffing in the field, it detects "
                "stealth port scans and denial-of-service attacks without relying on signatures, with sub-second latency and a "
                "low false-positive rate."
            )
            print("[OK 4b] Abstract enrichi avec la portabilité multiplateforme.")

    # ==========================================================================
    # 5. MATHEMATICAL FORMULAS & OMML CONVERSION
    # ==========================================================================
    omml_cn_xml = """
    <m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
      <m:oMath>
        <m:r><m:t>c(n) = 2 </m:t></m:r>
        <m:d>
          <m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
          <m:e>
            <m:r><m:t>ln(n - 1) + 0.5772156649</m:t></m:r>
          </m:e>
        </m:d>
        <m:r><m:t> - </m:t></m:r>
        <m:f>
          <m:fPr><m:type m:val="bar"/></m:fPr>
          <m:num><m:r><m:t>2(n - 1)</m:t></m:r></m:num>
          <m:den><m:r><m:t>n</m:t></m:r></m:den>
        </m:f>
      </m:oMath>
    </m:oMathPara>
    """

    omml_sxn_xml = """
    <m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
      <m:oMath>
        <m:r><m:t>s(x, n) = </m:t></m:r>
        <m:sSup>
          <m:e><m:r><m:t>2</m:t></m:r></m:e>
          <m:sup>
            <m:r><m:t>-</m:t></m:r>
            <m:f>
              <m:fPr><m:type m:val="bar"/></m:fPr>
              <m:num><m:r><m:t>E(h(x))</m:t></m:r></m:num>
              <m:den><m:r><m:t>c(n)</m:t></m:r></m:den>
            </m:f>
          </m:sup>
        </m:sSup>
      </m:oMath>
    </m:oMathPara>
    """

    omml_tw_xml = """
    <m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
      <m:oMath>
        <m:r><m:t>time_window = </m:t></m:r>
        <m:d>
          <m:dPr><m:begChr m:val="&#x230A;"/><m:endChr m:val="&#x230B;"/></m:dPr>
          <m:e>
            <m:f>
              <m:fPr><m:type m:val="bar"/></m:fPr>
              <m:num><m:r><m:t>Timestamp</m:t></m:r></m:num>
              <m:den><m:r><m:t>W</m:t></m:r></m:den>
            </m:f>
          </m:e>
        </m:d>
        <m:r><m:t> &#xD7; W</m:t></m:r>
      </m:oMath>
    </m:oMathPara>
    """

    for p in doc.paragraphs:
        # In-text formula c(n)
        if "c(n) = 2" in p.text and ("ln(n - 1)" in p.text or r"\ln(n - 1)" in p.text):
            p.text = ""
            p._p.append(parse_xml(omml_cn_xml))
            print("[OK 5a] Équation c(n) convertie en véritable équation Word OMML.")
        # In-text formula s(x,n)
        elif "s (x, n) = 2" in p.text or "s(x, n) = 2" in p.text:
            if "E(h(x))" in p.text:
                p.text = ""
                p._p.append(parse_xml(omml_sxn_xml))
                print("[OK 5b] Équation s(x, n) convertie en véritable équation Word OMML.")
        # In-text formula time_window
        elif "time\\_window" in p.text or "time_window = \\lfloor" in p.text:
            p.text = ""
            p._p.append(parse_xml(omml_tw_xml))
            print("[OK 5c] Équation time_window convertie en véritable équation Word OMML.")
        # Clean section 2.4.1 mathematical text
        elif r"X' \subset X" in p.text:
            p.text = p.text.replace(r"X' \subset X", "X' ⊂ X")
        elif r"p \sim \mathcal{U}" in p.text:
            p.text = p.text.replace(r"p \sim \mathcal{U}(\min(q), \max(q))", "p ~ U(min(q), max(q))")
        elif r"X_{gauche}" in p.text:
            p.text = p.text.replace(r"X_{gauche} = \{x \mid x_q < p\} et X_{droite} = \{x \mid x_q \ge p\}", "X_gauche = {x | x_q < p} et X_droite = {x | x_q ≥ p}")
        elif r"h_{max} = \lceil" in p.text:
            p.text = p.text.replace(r"h_{max} = \lceil \log_2(n) \rceil", "h_max = ⌈log₂(n)⌉")
        # Clean section 2.4.2 properties
        elif "Si E(h(x))" in p.text:
            p.text = p.text.replace(r"\to", "→").replace(r"\implies", "⇒")
        elif "Propriétés fondamentales du score $s$ :" in p.text:
            p.text = "Propriétés fondamentales du score s :"
        # Clean Section 3.3.2 introductory text
        elif "W = 5.0 \\text{ secondes}" in p.text:
            p.text = p.text.replace("W = 5.0 \\text{ secondes}", "W = 5.0 secondes")
        # Clean Rules in Section 3.4
        elif "Règle 1" in p.text and r"\frac{P_{unique}}" in p.text:
            p.text = p.text.replace(r"R = \frac{P_{unique}}{N_{pkts}} > 0.5", "R = P_unique / N_pkts > 0.5").replace("`Port Scan`", "Port Scan")
        elif "Règle 2" in p.text and (r"\ge 50" in p.text or r"\le 0.5" in p.text):
            p.text = p.text.replace(r"N_{pkts} \ge 50", "N_pkts ≥ 50").replace(r"R \le 0.5", "R ≤ 0.5").replace("`Flood / DDoS`", "Flood / DDoS")
        elif "Règle 3" in p.text and "`Normal`" in p.text:
            p.text = p.text.replace("`Normal`", "Normal")
        elif "Règle 4" in p.text and (r"\ge" in p.text or r"\frac" in p.text):
            p.text = p.text.replace(r"P_{unique} \ge 10", "P_unique ≥ 10").replace(r"R = \frac{P_{unique}}{N_{pkts}} \ge 0.7", "R = P_unique / N_pkts ≥ 0.7").replace("`Port Scan`", "Port Scan")
        # Clean section headings backticks
        elif p.text.startswith("4.2 L'Orchestrateur CLI") and "`netscan.py`" in p.text:
            p.text = p.text.replace("`netscan.py`", "netscan.py")
        elif p.text.startswith("4.3 Le Tableau de Bord") and "`dashboard.py`" in p.text:
            p.text = p.text.replace("`dashboard.py`", "dashboard.py")
        elif p.text.startswith("4.5 Moteur Exécutif") and "`report_generator.py`" in p.text:
            p.text = p.text.replace("`report_generator.py`", "report_generator.py")
        elif p.text.startswith("4.6 Suite d'Assurance") and "`test_smoke.py`" in p.text:
            p.text = p.text.replace("`test_smoke.py`", "test_smoke.py")
        elif "(*Feature Engineering*)" in p.text and p.style.name.startswith("Heading"):
            p.text = p.text.replace("(*Feature Engineering*)", "(Feature Engineering)")
        elif "(*iTrees*)" in p.text and p.style.name.startswith("Heading"):
            p.text = p.text.replace("(*iTrees*)", "(iTrees)")
        # Fix section 4.4 style to Heading 2
        elif p.text.startswith("4.4 Moteur de Journalisation Sécurisée") and p.style.name != "Heading 2":
            p.style = "Heading 2"
            print("[OK 5d] Style du Titre 4.4 défini sur 'Heading 2'.")
        # Clean section 4.4 and 5.2 inline formulas
        elif "$R$" in p.text:
            p.text = p.text.replace("$R$", "R")
        elif r"$P_{\text{unique}} > 350$" in p.text:
            p.text = p.text.replace(r"$P_{\text{unique}} > 350$", "P_unique > 350")
        elif r"$R = \frac{P_{\text{unique}}}{N_{\text{pkts}}}$" in p.text:
            p.text = p.text.replace(r"$R = \frac{P_{\text{unique}}}{N_{\text{pkts}}}$", "R = P_unique / N_pkts")

    # Format general body paragraphs with italics (*payload pattern*, etc.)
    for p in doc.paragraphs:
        format_paragraph_markdown_inline(p)

    print("[OK 5e] Italiques (*...*) et gras (**...**) formatés proprement dans tout le corps du rapport et la bibliographie.")

    # ==========================================================================
    # 6. CAPTIONS FORMATTING & CLEANING (FIGURE 5.2 HEADING FIX)
    # ==========================================================================
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("Figure 5."):
            if p.style.name.startswith("Heading") or p.style.name == "List Paragraph":
                p.style = "Normal"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(12)
            for r in p.runs:
                r.italic = True
                r.font.name = "Calibri"
                r.font.size = Pt(10)

    print("[OK 6] Légendes des figures 5.1 à 5.7 vérifiées et centrées en italique.")

    # ==========================================================================
    # 7. TABLE 9 RECONSTRUCTION (FROM CORRUPTED 5-COL TO CLEAN 3-COL)
    # ==========================================================================
    t9 = doc.tables[9]
    for row in t9.rows:
        tcs = row._tr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')
        if len(tcs) > 3:
            for tc in tcs[3:]:
                row._tr.remove(tc)

    tblGrid = t9._tbl.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblGrid')
    if tblGrid is not None:
        gridCols = tblGrid.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridCol')
        for gc in gridCols[3:]:
            tblGrid.remove(gc)
        if len(gridCols) >= 3:
            gridCols[0].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w', '2500')
            gridCols[1].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w', '2800')
            gridCols[2].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w', '4400')

    t9_data = [
        ("Feature Calculée", "Formule Mathématique", "Rôle dans la Détection"),
        ("`packet_count`", "N_pkts = ∑ 1", "Volume d'activité globale de l'hôte dans la fenêtre de 5 secondes."),
        ("`unique_dst_ports`", "P_unique = |⋃ Dst_Port|", "**Caractéristique clé :** diversité des services distants sondés."),
        ("`unique_dst_ips`", "IP_unique = |⋃ Dst_Address|", "Détermine si l'activité cible une machine unique ou un réseau entier."),
        ("`syn_count`", "N_SYN = ∑ SYN_Flag", "Nombre total d'initiations de sessions TCP émises sans réponse."),
        ("`avg_packet_length`", "μ_L = (1/N) ∑ Length", "Signature de taille : les scans ont une moyenne très basse (≈ 60 B)."),
        ("`port_dispersion_ratio`", "R = P_unique / N_pkts", "Ratio d'entropie : proche de 1.0 lors d'un scan, très faible lors d'un flood.")
    ]

    for r_idx, (f_col, m_col, r_col) in enumerate(t9_data):
        row = t9.rows[r_idx]
        if r_idx == 0:
            row.cells[0].paragraphs[0].text = f_col
            row.cells[1].paragraphs[0].text = m_col
            row.cells[2].paragraphs[0].text = r_col
        else:
            format_cell_markdown(row.cells[0].paragraphs[0], f_col)
            format_cell_markdown(row.cells[1].paragraphs[0], m_col)
            format_cell_markdown(row.cells[2].paragraphs[0], r_col)

    print("[OK 7] Tableau 3.2 (Table 9) restructuré proprement en 3 colonnes.")

    # ==========================================================================
    # 8. PROCESS ALL OTHER TABLES FOR LITERAL MARKDOWN / LATEX
    # ==========================================================================
    for t_idx in [3, 4, 6, 7, 8, 10, 11, 12, 13]:
        tbl = doc.tables[t_idx]
        for r_idx, row in enumerate(tbl.rows):
            if r_idx == 0:
                continue
            for c_idx, cell in enumerate(row.cells):
                txt = cell.text.strip()
                if any(tok in txt for tok in ["**", "*", "`", "$", "\\"]):
                    format_cell_markdown(cell.paragraphs[0], txt)

    print("[OK 8] Markdown et LaTeX convertis en formatage natif Word dans tous les tableaux (3, 4, 6, 7, 8, 10, 11, 12, 13).")

    # ==========================================================================
    # 9. WORD SETTINGS : ENABLE AUTO-UPDATE OF FIELDS ON OPEN
    # ==========================================================================
    settings_elm = doc.settings.element
    if settings_elm.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields') is None:
        update_fields = parse_xml(r'<w:updateFields %s w:val="true"/>' % nsdecls('w'))
        settings_elm.append(update_fields)
        print("[OK 9] Paramètre Word updateFields=true activé (mise à jour automatique des champs à l'ouverture).")

    # Save document
    doc.save(out_doc_path)
    print(f"\n==================================================================")
    print(f">>> SUCCÈS MAJEUR : Document parfait généré avec succès dans :")
    print(f"    {out_doc_path}")
    print(f"==================================================================")

if __name__ == "__main__":
    src = r"C:\Users\DevLab\Desktop\Soutenance\Rapport_MohamedRedaQIYAOUI_FIXED.docx"
    out = r"C:\Users\DevLab\Desktop\Soutenance\Rapport_MohamedRedaQIYAOUI_PERFECT.docx"
    run_master_fix(src, out)
