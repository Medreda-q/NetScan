import docx

doc = docx.Document(r"C:\Users\DevLab\Desktop\Soutenance\Rapport_MohamedRedaQIYAOUI_FIXED.docx")

toc_entries = []
for i in range(58, 107):
    p = doc.paragraphs[i]
    toc_entries.append((i, p.text.strip()))

body_headings = []
for i, p in enumerate(doc.paragraphs):
    if i > 110:
        txt = p.text.strip()
        if p.style.name.startswith("Heading") or p.style.name.startswith("Titre") or any(txt.startswith(f"{x}.") or txt.startswith(f"{x} ") for x in range(1, 10)):
            body_headings.append((i, p.style.name, txt))

print("=== TOC ENTRIES ===")
for idx, txt in toc_entries:
    print(f"ToC[{idx}]: {txt}")

print("\n=== BODY HEADINGS ===")
for idx, style, txt in body_headings:
    print(f"Body[{idx}] [{style}]: {txt}")
