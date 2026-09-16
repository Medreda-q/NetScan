import docx

doc = docx.Document(r"C:\Users\DevLab\Desktop\Soutenance\Avec Commune\Rapport_MohamedRedaQIYAOUI.docx")

print("=== IMAGES IN DOCUMENT ===")
rels = doc.part.rels
img_count = 0
for rel in rels.values():
    if "image" in rel.target_ref:
        img_count += 1
        print(f"Image {img_count}: {rel.target_ref}")

print(f"Total embedded images: {img_count}")

print("\n=== HEADINGS AND LEVELS ===")
for i, p in enumerate(doc.paragraphs):
    if p.style.name.startswith("Heading"):
        print(f"[{p.style.name}] {p.text}")

print("\n=== TABLES SUMMARY ===")
for t_idx, table in enumerate(doc.tables):
    print(f"Table {t_idx+1} ({len(table.rows)} rows, {len(table.columns)} cols): Header = {[c.text.strip() for c in table.rows[0].cells]}")
