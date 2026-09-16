import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
tabs = parse_xml(r'<w:tabs %s><w:tab w:val="right" w:leader="dot" w:pos="9350"/></w:tabs>' % nsdecls('w'))
pPr.append(tabs)

r1 = p.add_run("2.  Chapitre I : Présentation du Cadre Académique")
r1.bold = True
r2 = p.add_run("\t12")
r2.bold = True

doc.save("test_tab.docx")
print("Saved test_tab.docx successfully!")
