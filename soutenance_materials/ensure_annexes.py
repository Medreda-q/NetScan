import sys
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding='utf-8')
doc_path = r"C:\Users\DevLab\Desktop\Soutenance\Avec Commune\Rapport_MohamedRedaQIYAOUI.docx"
doc = docx.Document(doc_path)

# Check if Annexes are present
has_annexes = any("ANNEXE" in p.text.upper() for p in doc.paragraphs)
print(f"Has Annexes: {has_annexes}")

if not has_annexes:
    print("Appending Annexes...")
    doc.add_page_break()
    p_annexe_h = doc.add_paragraph("ANNEXES", style="Heading 1")

    p_annexe_a = doc.add_paragraph("Annexe A : Fiche Technique d'Architecture & Modules Logiciels NetScan", style="Heading 2")
    p_annexe_intro = doc.add_paragraph(
        "Le tableau ci-dessous synthétise la cartographie modulaire de NetScan telle que présentée aux membres du jury. "
        "Il détaille le rôle opérationnel de chaque fichier, ses flux d'entrée/sortie et les commandes d'exécution correspondantes."
    )
    p_annexe_intro.paragraph_format.line_spacing = 1.5

    table_data = [
        ("Module / Fichier", "Rôle & Mission Principale", "Entrées / Sources", "Sorties / Livrables", "Commande d'Exécution"),
        ("netscan.py", "Orchestrateur principal & CLI", "Arguments CLI, flux réseau brut", "Logs console ANSI, initialisation", "python netscan.py --help"),
        ("detection_engine.py", "Moteur IA Dual-Layer (Isolation Forest)", "Flux direct (Npcap/libpcap) ou PCAP", "Scores d'anomalie [0.0 - 1.0], alertes", "python netscan.py --pcap <capture.pcap>"),
        ("dashboard.py", "Tableau de bord SOC interactif", "Métriques streaming, logs CSV", "Interface visuelle Streamlit temps réel", "python netscan.py --dashboard"),
        ("alerting.py", "Alerting instantané & défense active", "Événements anormaux (score >= 0.60)", "Webhooks Discord (< 300 ms), règles iptables", "Automatique en cours d'analyse"),
        ("report_generator.py", "Générateur exécutif de rapports forensiques", "Métriques d'incident, timeline", "Rapports HTML5 et PDF (NIST SP 800-61)", "python netscan.py --pcap <file> --report"),
        ("test_smoke.py", "Banc de tests et assurance qualité", "Environnement local Python", "Validation des 5 modules en < 3 secondes", "python test_smoke.py"),
    ]

    table = doc.add_table(rows=len(table_data), cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    col_widths = [Cm(2.6), Cm(4.1), Cm(3.2), Cm(3.2), Cm(3.4)]

    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = col_widths[c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            if r_idx == 0:
                run.font.size = Pt(9.5)
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0A2540"/>')
                cell._tc.get_or_add_tcPr().append(shading)
            else:
                run.font.size = Pt(8.5)
                if c_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(10, 37, 64)
                if r_idx % 2 == 1:
                    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
                    cell._tc.get_or_add_tcPr().append(shading)

    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

    p_tbl_cap = doc.add_paragraph("Tableau A.1 : Matrice descriptive des composants logiciels du système NetScan.", style="Normal")
    p_tbl_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tbl_cap.paragraph_format.space_after = Pt(14)
    for r in p_tbl_cap.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)
        r.font.italic = True
        r.font.color.rgb = RGBColor(90, 90, 90)

    p_annexe_b = doc.add_paragraph("Annexe B : Guide d'Exécution et Traces Opérationnelles", style="Heading 2")
    p_annexe_b_txt = doc.add_paragraph(
        "NetScan prend en charge deux modes d'exploitation distincts, adaptés respectivement à l'investigation forensique "
        "post-mortem et à la détection proactive en environnement de production :\n\n"
        "1. Mode Forensique Hors-Ligne (Rejeu de capture PCAP) :\n"
        "   python netscan.py --pcap sample_nmap_attack.pcap --report\n"
        "   Ce mode lit séquentiellement les trames, extrait les vecteurs 5D, applique la prédiction Isolation Forest "
        "et génère un rapport forensique d'incident complet conforme aux directives NIST SP 800-61.\n\n"
        "2. Mode Surveillance Temps Réel (Calibrate-Then-Monitor) :\n"
        "   python netscan.py --live --interface eth0 --calibrate 120 --window 5\n"
        "   Ce mode observe le trafic normal pendant 120 secondes pour calibrer la ligne de base, puis bascule en surveillance "
        "continue par micro-fenêtres de 5 secondes avec notification Discord et génération automatique de règles de blocage iptables."
    )
    p_annexe_b_txt.paragraph_format.line_spacing = 1.5
    for r in p_annexe_b_txt.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    doc.save(doc_path)
    print("Annexes successfully appended and saved.")
else:
    print("Annexes already present.")
