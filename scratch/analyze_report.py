import docx
import os
import sys

doc_path = r"C:\Users\DevLab\Documents\Rapport_NetScan_corrige (1).docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Dump full outline and image locations
with open("scratch/report_analysis.txt", "w", encoding="utf-8") as out:
    for i, p in enumerate(doc.paragraphs):
        xml = p._element.xml
        has_drawing = "w:drawing" in xml
        txt = p.text.strip()
        style = p.style.name if p.style else ""
        
        flags = []
        if has_drawing:
            flags.append("IMG")
        if style:
            flags.append(f"style={style}")
            
        flag_str = f" [{', '.join(flags)}]" if flags else ""
        
        if has_drawing or txt.startswith("Figure ") or txt.startswith("Tableau ") or "Heading" in style or "Titre" in style or (len(txt) < 80 and any(txt.startswith(x) for x in ["Chapitre", "1.", "2.", "3.", "4.", "5.", "6.", "Conclusion", "Introduction", "Liste"])):
            out.write(f"P{i:04d}{flag_str}: {txt}\n")

print("Report outline and images written to scratch/report_analysis.txt")
