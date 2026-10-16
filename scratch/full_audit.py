import docx

doc = docx.Document(r"C:\Users\DevLab\Documents\Rapport_NetScan_corrige (1).docx")

with open("scratch/full_report_audit.txt", "w", encoding="utf-8") as out:
    out.write("=== 1. TABLES AUDIT ===\n")
    out.write(f"Total tables found in document: {len(doc.tables)}\n")
    for t_idx, table in enumerate(doc.tables):
        rows = len(table.rows)
        cols = len(table.columns)
        first_cell = table.cell(0, 0).text.strip().replace("\n", " ") if rows > 0 and cols > 0 else ""
        out.write(f"Table {t_idx:02d}: {rows} rows x {cols} cols | Header/First cell: '{first_cell[:80]}'\n")

    out.write("\n=== 2. LISTE DES TABLEAUX (FRONT MATTER) ===\n")
    for i in range(134, 150):
        txt = doc.paragraphs[i].text.strip()
        if txt.startswith("Tableau "):
            out.write(f"P{i:04d}: {txt}\n")

    out.write("\n=== 3. TABLEAU CAPTIONS IN BODY ===\n")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if i >= 150 and txt.startswith("Tableau "):
            out.write(f"P{i:04d}: {txt}\n")

    out.write("\n=== 4. HEADINGS IN BODY ===\n")
    for i, p in enumerate(doc.paragraphs):
        style = p.style.name if p.style else ""
        txt = p.text.strip()
        if ("Heading" in style or "Titre" in style) and i >= 190:
            out.write(f"P{i:04d} [{style}]: {txt}\n")

    out.write("\n=== 5. TABLE DES MATIERES (FRONT MATTER) ===\n")
    for i in range(58, 110):
        txt = doc.paragraphs[i].text.strip()
        if txt:
            out.write(f"P{i:04d}: {txt}\n")

print("Full audit written to scratch/full_report_audit.txt")
