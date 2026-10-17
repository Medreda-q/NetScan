import docx

doc = docx.Document(r"C:\Users\DevLab\Documents\Rapport_NetScan_corrige (1).docx")
print("P0012 XML:")
print(doc.paragraphs[12]._element.xml)
print("\nP0013 XML:")
print(doc.paragraphs[13]._element.xml)
print("\nP0014 XML:")
print(doc.paragraphs[14]._element.xml)
