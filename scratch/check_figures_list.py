import docx

doc = docx.Document(r"C:\Users\DevLab\Documents\Rapport_NetScan_corrige (1).docx")

# Find all occurrences of Figure in text
print("=== LIST OF ALL FIGURE CAPTIONS IN TEXT ===")
captions_in_text = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("Figure ") or txt.startswith("Fig. "):
        captions_in_text.append((i, txt))
        print(f"P{i:04d}: {txt}")

print("\n=== LIST OF FIGURES IN FRONT MATTER (TABLE OF FIGURES) ===")
for i in range(110, 134):
    txt = doc.paragraphs[i].text.strip()
    if txt.startswith("Figure "):
        print(f"List P{i:04d}: {txt}")
