import docx

doc = docx.Document(r"C:\Users\DevLab\Documents\Rapport_NetScan_corrige (1).docx")

with open("scratch/chap6_text.txt", "w", encoding="utf-8") as out:
    for i in range(460, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        style = p.style.name if p.style else ""
        txt = p.text.strip()
        if txt:
            out.write(f"P{i:04d} [{style}]: {txt}\n")

print("Dumped Chapter 6 to scratch/chap6_text.txt")
