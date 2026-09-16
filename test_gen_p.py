import re
import docx

def format_general_paragraph(p):
    # Skip if contains math equations or drawing/images
    xml_str = p._p.xml
    if "oMath" in xml_str or "blip" in xml_str:
        return
    
    txt = p.text
    if "*" not in txt and "`" not in txt and "$" not in txt:
        return
    
    # Specific cleanups
    txt = txt.replace("$R$", "R").replace("$< 30$", "< 30")
    
    # Tokenize on **bold**, *italic*, `code`
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    parts = pattern.split(txt)
    if len(parts) <= 1:
        return

    # Clear runs while preserving paragraph style and alignment
    p.text = ""
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            p.add_run(part)

print("Paragraph formatting function ready!")
